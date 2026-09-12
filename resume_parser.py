"""Module 1 & 2 (extraction part): read resume text out of PDF/DOCX files."""

from __future__ import annotations

import io
import os

from pypdf import PdfReader
from docx import Document

MAX_FILE_SIZE_MB = 5
ALLOWED_EXTENSIONS = {".pdf", ".docx"}


class ResumeParseError(Exception):
    """Raised when a resume file cannot be validated or read."""


def validate_file(filename: str, size_bytes: int) -> None:
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeParseError(
            f"Unsupported file type '{ext}'. Please upload a PDF or DOCX resume."
        )
    if size_bytes > MAX_FILE_SIZE_MB * 1024 * 1024:
        raise ResumeParseError(
            f"File is too large ({size_bytes / (1024 * 1024):.1f} MB). "
            f"Maximum allowed size is {MAX_FILE_SIZE_MB} MB."
        )


def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages_text = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages_text)


def extract_text_from_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Dispatch extraction based on file extension. Raises ResumeParseError on failure."""
    validate_file(filename, len(file_bytes))
    ext = os.path.splitext(filename)[1].lower()
    try:
        if ext == ".pdf":
            text = extract_text_from_pdf(file_bytes)
        else:
            text = extract_text_from_docx(file_bytes)
    except Exception as exc:  # noqa: BLE001 - surface any parser failure uniformly
        raise ResumeParseError(f"Could not read '{filename}': {exc}") from exc

    if not text.strip():
        raise ResumeParseError(
            f"No readable text found in '{filename}'. The file may be a scanned image."
        )
    return text
